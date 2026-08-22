#!/usr/bin/env python3
"""Zapis wyników uruchomienia do .xlsx (wywoływane przez uruchom-lokalnie.js)."""
import json
import sys
import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

STALA_DATA = datetime.datetime(2026, 8, 22, 9, 0, 0)
KRAWEDZ = Border(*[Side(style='thin', color='BFBFBF')] * 4)
SZARY = PatternFill('solid', fgColor='EFEFEF')
ZOLTY = PatternFill('solid', fgColor='FFF2CC')


def zapisz(wpis, katalog):
    wb = Workbook()
    wb.remove(wb.active)
    for ark in wpis['arkusze']:
        ws = wb.create_sheet(ark['nazwa'][:31])
        for wiersz in ark['wiersze']:
            ws.append(wiersz)
        ws.column_dimensions['A'].width = 26
        ws.column_dimensions['B'].width = 34
        for kol in 'CDEFGHIJKLM':
            ws.column_dimensions[kol].width = 7
        for w in ws.iter_rows():
            for c in w:
                if c.column > 2:
                    c.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    c.alignment = Alignment(vertical='center', wrap_text=True)
                if c.value not in (None, ''):
                    c.border = KRAWEDZ
        for w in ws.iter_rows(min_col=1, max_col=1):
            v = str(w[0].value or '')
            if v.startswith('zajęcia objęte'):
                for c in ws[w[0].row]:
                    c.fill = SZARY
                    c.font = Font(bold=True, italic=True)
            elif v.startswith('UWAGI'):
                for c in ws[w[0].row]:
                    c.fill = ZOLTY
                    c.font = Font(bold=True)
        for c in ws[1] + ws[4] + ws[5]:
            c.font = Font(bold=True)
        ws.freeze_panes = 'C6'
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.created = STALA_DATA
    wb.properties.modified = STALA_DATA
    wb.properties.creator = 'generator wniosków ALA'
    sciezka = '%s/%s.xlsx' % (katalog, wpis['nazwa'])
    wb.save(sciezka)
    return sciezka


if __name__ == '__main__':
    dane = json.load(open(sys.argv[1], encoding='utf8'))
    for wpis in dane:
        print('   zapisano ' + zapisz(wpis, sys.argv[2]).split('/')[-1])
